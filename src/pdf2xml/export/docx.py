"""DOCX writer (python-docx): catalogue styles → named Word styles, runs → formatted runs.

Mapping: ``body`` → Normal, ``hN`` → Heading N, ``caption`` → Caption, ``code`` → "Code",
anything else → a custom paragraph style "PX <key>" based on Normal.
Phase 1 limits: list markers are written literally with a hanging indent (real Word numbering
is Phase 2); hyperlinks are plain text; formulas are text (OMML is Phase 3).
"""

from __future__ import annotations

import contextlib
from pathlib import Path
from typing import Any

from docx import Document as new_docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from pdf2xml.merge.text import CID_RE
from pdf2xml.model import (
    Block,
    Box,
    Caption,
    Code,
    Document,
    Figure,
    Formula,
    Heading,
    ListBlock,
    ListItem,
    Paragraph,
    ParaProps,
    Run,
    Style,
    Table,
)

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_LIST_INDENT = 18.0   # pt per nesting level
_HANGING = 18.0       # pt between marker and text


def _set_font_name(rpr_owner: Any, name: str) -> None:
    """Set a font on a run/style and drop theme fonts, which would otherwise win."""
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


class DocxWriter:
    def __init__(self, doc: Document, out_dir: Path) -> None:
        self.doc = doc
        self.out_dir = out_dir
        self.d = new_docx()
        self.style_names: dict[str, str] = {}
        self.text_width = Pt(450)

    # -------------------------------------------------------------- setup

    def _page_setup(self) -> None:
        if not self.doc.pages:
            return
        p1 = self.doc.pages[0]
        boxes = [b.bbox for b in self.doc.body if b.page == p1.n]
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Pt(p1.width), Pt(p1.height)
        if boxes:
            left = max(min(b[0] for b in boxes), 21.6)
            right = max(p1.width - max(b[2] for b in boxes), 21.6)
            top = max(min(b[1] for b in boxes), 21.6)
            bottom = max(p1.height - max(b[3] for b in boxes), 21.6)
            sec.left_margin, sec.right_margin = Pt(left), Pt(right)
            sec.top_margin, sec.bottom_margin = Pt(top), Pt(bottom)
            self.text_width = Pt(p1.width - left - right)

    def _apply_style(self, st: Any, s: Style) -> None:
        font = st.font
        _set_font_name(st.element, s.font)
        font.size = Pt(s.size)
        font.bold = s.bold
        font.italic = s.italic
        font.small_caps = s.small_caps
        font.color.rgb = _rgb(s.color)
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)

    def _styles(self) -> None:
        styles = self.d.styles
        for key, s in sorted(self.doc.styles.items()):
            if key == "body":
                name = "Normal"
            elif s.role == "heading" and s.level and key == f"h{s.level}":
                name = f"Heading {s.level}"
            elif key == "caption":
                name = "Caption"
            elif key == "code":
                name = "Code"
            else:
                name = f"PX {key}"
            try:
                st = styles[name]
            except KeyError:
                st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = styles["Normal"]
            self._apply_style(st, s)
            self.style_names[key] = name
        if "Normal" not in self.style_names.values():
            normal = styles["Normal"].paragraph_format
            normal.space_after = Pt(0)

    # -------------------------------------------------------------- content

    def _para_format(self, p: Any, props: ParaProps | None, size: float) -> None:
        if props is None:
            return
        pf = p.paragraph_format
        if props.align:
            pf.alignment = _ALIGN[props.align]
        if props.indent_first is not None:
            pf.first_line_indent = Pt(props.indent_first)
        if props.space_before is not None:
            pf.space_before = Pt(props.space_before)
        if props.line_height:
            # Word's "single" is ~1.15× the font size for typical fonts.
            pf.line_spacing = max(0.8, min(3.0, round(props.line_height / 1.15, 2)))

    def _runs(self, p: Any, runs: list[Run]) -> None:
        for r in runs:
            run = p.add_run(r.text)
            f = run.font
            if r.bold is not None:
                f.bold = r.bold
            if r.italic is not None:
                f.italic = r.italic
            if r.underline is not None:
                f.underline = r.underline
            if r.strike is not None:
                f.strike = r.strike
            if r.sup:
                f.superscript = True
            if r.sub:
                f.subscript = True
            if r.small_caps is not None:
                f.small_caps = r.small_caps
            if r.color:
                f.color.rgb = _rgb(r.color)
            if r.size is not None:
                f.size = Pt(r.size)
            if r.font:
                _set_font_name(run._r, r.font)

    def _style_of(self, b: Block, default: str = "Normal") -> str:
        return self.style_names.get(b.style or "", default)

    def _size_of(self, b: Block) -> float:
        s = self.doc.styles.get(b.style or "") or self.doc.styles.get("body")
        return s.size if s else 10.0

    def _block(self, b: Block, depth: int = 0) -> None:
        if isinstance(b, Heading):
            name = self.style_names.get(b.style or "", f"Heading {b.level}")
            p = self.d.add_paragraph(style=name)
            self._para_format(p, b.para, self._size_of(b))
            self._runs(p, b.runs)
        elif isinstance(b, Paragraph | Caption):
            p = self.d.add_paragraph(style=self._style_of(b))
            self._para_format(p, b.para, self._size_of(b))
            self._runs(p, b.runs)
        elif isinstance(b, Code):
            p = self.d.add_paragraph(style=self._style_of(b))
            self._runs(p, b.runs)
        elif isinstance(b, Formula):
            p = self.d.add_paragraph(style=self._style_of(b))
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._runs(p, b.runs or [Run(text=b.latex or "")])
        elif isinstance(b, ListBlock):
            for item in b.items:
                self._list_item(item, depth)
        elif isinstance(b, Table):
            self._table(b)
        elif isinstance(b, Figure):
            self._figure(b)
        elif isinstance(b, Box):
            for child in b.children:
                self._block(child, depth)
        elif isinstance(b, ListItem):
            self._list_item(b, depth)

    def _list_item(self, item: ListItem, depth: int) -> None:
        p = self.d.add_paragraph(style=self._style_of(item))
        pf = p.paragraph_format
        pf.left_indent = Pt(_LIST_INDENT * depth + _HANGING)
        pf.first_line_indent = Pt(-_HANGING)
        pf.tab_stops.add_tab_stop(Pt(_LIST_INDENT * depth + _HANGING))
        if item.para and item.para.space_before is not None:
            pf.space_before = Pt(item.para.space_before)
        marker = item.marker or "•"
        if CID_RE.fullmatch(marker):   # unmappable bullet glyph kept verbatim in the canonical data
            marker = "•"
        p.add_run(f"{marker}\t")
        self._runs(p, item.runs)
        for child in item.children:
            self._block(child, depth + 1)

    def _table(self, t: Table) -> None:
        n_rows = len(t.rows)
        if n_rows == 0:
            return
        table = self.d.add_table(rows=n_rows, cols=max(1, t.n_cols))
        with contextlib.suppress(KeyError):
            table.style = self.d.styles["Table Grid"]
        occupied: set[tuple[int, int]] = set()
        for r, row in enumerate(t.rows):
            c = 0
            for cell in row:
                while (r, c) in occupied:
                    c += 1
                if c >= t.n_cols:
                    break
                r2 = min(r + cell.rowspan - 1, n_rows - 1)
                c2 = min(c + cell.colspan - 1, t.n_cols - 1)
                for rr in range(r, r2 + 1):
                    for cc in range(c, c2 + 1):
                        occupied.add((rr, cc))
                target = table.cell(r, c)
                if (r2, c2) != (r, c):
                    target = target.merge(table.cell(r2, c2))
                p = target.paragraphs[0]
                self._runs(p, cell.runs)
                if cell.header:
                    for run in p.runs:
                        if run.font.bold is None:
                            run.font.bold = True
                c = c2 + 1

    def _figure(self, f: Figure) -> None:
        if not f.image or not (self.out_dir / f.image).exists():
            return
        width = min(Pt(f.bbox[2] - f.bbox[0]), self.text_width)
        p = self.d.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(self.out_dir / f.image), width=width)

    # -------------------------------------------------------------- entry

    def write(self, path: Path) -> Path:
        self._page_setup()
        self._styles()
        if self.doc.meta.title:
            self.d.core_properties.title = self.doc.meta.title
        for b in self.doc.body:
            self._block(b)
        for b in [*self.doc.footnotes, *self.doc.references]:
            self._block(b)
        self.d.save(str(path))
        return path


def write_docx(doc: Document, out_dir: Path, path: Path) -> Path:
    return DocxWriter(doc, out_dir).write(path)
